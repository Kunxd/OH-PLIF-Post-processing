#!/usr/bin/env python3
"""
run_wabel_batch.py  (v8 — batch / resume / parallel)
Based on run_wabel_canny.py v7 (Wabel_binarization_cmap_generation):
  Wabel local thresholding → binarize each frame → per-position c̄ averaging
  → ray-scan c̄ contours → zero-fill pos0 + linear extension to nozzle lip
  → surface area by revolution → flame height (topmost contour point).
  The image-processing maths is unchanged from v7.

New in v8
  • Recursive input search, flame name = everything before "_pos<N>"
  • One global worker pool, 1 job = 1 position (all flames queued at once)
  • Chunked .spe reading (RAM per worker ~0.3–0.5 GB instead of ~10 GB)
  • Resume: positions with a valid done.json are skipped
  • Toggles DO_CBAR / DO_INSTANT (independent)
  • Instantaneous composite: best-match frame chaining (dynamic programming)
  • Per-flame: "average composites" folder, c-lines .npz, Excel, settings log
  • Final ALL_conditions_results.xlsx over all flames
  • Run log file + warning summary

Requires spe_reader.py in the same folder.  Python ≥ 3.8.
"""
import os
# One thread per worker process — must be set BEFORE numpy / cv2 import,
# otherwise 4 workers × 8 library threads fight over 8 CPU threads.
for _v in ("OMP_NUM_THREADS", "OPENBLAS_NUM_THREADS", "MKL_NUM_THREADS",
           "NUMEXPR_NUM_THREADS", "VECLIB_MAXIMUM_THREADS"):
    os.environ.setdefault(_v, "1")

import sys, re, json, time, shutil, hashlib, traceback, datetime, glob, math
from collections import OrderedDict
import numpy as np
from scipy.ndimage import maximum_filter
import cv2
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.lines import Line2D
from matplotlib.colors import LinearSegmentedColormap

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from spe_reader import read_spe_metadata, _first_real_frame_index

# ╔════════════════════════════════════════════════════════════════════╗
# ║                       CONFIG — EDIT HERE                           ║
# ╚════════════════════════════════════════════════════════════════════╝
# Searched recursively (all subfolders).  Background files: "bg" after posN.
INPUT_DIR       = "/media/USERNAME/EXTERNAL_DRIVE/scaling_spe/OHLIF"
OUTDIR          = "/media/USERNAME/EXTERNAL_DRIVE/Wabel_results/batch_v8"

# ── What to compute (independent toggles) ───────────────────────────
DO_CBAR         = True   # c̄ averaging → c-lines → A_F, h  (all frames)
DO_INSTANT      = True   # instantaneous fronts → stitched composite

# ── Parallelism / IO ────────────────────────────────────────────────
NUM_WORKERS     = 4      # i7-11th gen 4C/8T laptop: 4.  1 = no multiprocessing
CHUNK_FRAMES    = 25     # frames read per chunk (RAM vs. disk-seek trade-off)

# ── Run control ─────────────────────────────────────────────────────
DRY_RUN                 = False  # True: only discover + resume scan, no processing
MAX_FLAMES              = None   # e.g. 1 → only the first flame (alphabetical); None = all
ONLY_FLAMES             = []     # e.g. ["Re15000_phi1"] → only flames whose name contains
                                 # one of these texts; [] = no filter
EXCLUDE_FLAMES          = []     # e.g. ["Re26992"] → skip flames whose name contains these
DUPLICATE_POLICY        = "average"  # same flame+position recorded more than once:
                                 # "average" = average all recordings
                                 # "newest"  = use only the last day-folder's recording
                                 # "oldest"  = use only the first day-folder's recording
REBUILD_FINISHED_FLAMES = False  # True: redo composites/Excel of finished flames
ALLOW_POSITION_GAPS     = False  # True: build composites even if e.g. pos3 missing
MISSING_BG_ACTION       = "warn" # "warn": process without bg subtraction (logged)
                                 # "skip": do not process that position
NOZZLE_DIAM_FALLBACK_MM = None   # used only if "D<n>" is not found in the name

# ── Frames ──────────────────────────────────────────────────────────
MAX_FRAMES      = None   # cap frames per position for c̄ (None = all)
EXPECTED_FRAMES = 1000   # warn if a position has a different number
N_INSTANT       = 50     # frames per position stored for instantaneous fronts
N_INSTANT_COMPOSITES = 3 # stitched instantaneous composites per flame
STITCH_BAND_PX  = 3      # rows at each position edge used for matching
INSTANT_AREA    = True   # revolve the stitched instantaneous front (180° per side,
                         # like the c-lines) → instantaneous surface area A_inst
GAP_CLOSE_PX    = 5      # seal gaps up to ~this size in the burnt region before
                         # separating inner front from the outer (ambient) edge
MIN_POCKET_AREA_PX = 50  # enclosed unburnt pockets smaller than this are ignored
FRONT_SMOOTH_PX = 5      # smoothing window [px] along the front when measuring its
                         # length (removes the pixel staircase)

# ── Image processing (unchanged from v7) ────────────────────────────
SPATIAL_RES     = 0.05
CROP_BOTTOM_MM  = 7.75
CROP_TOP_MM     = 1.0
NOISE_FLOOR     = 10
WINDOW_SIZE     = 51
THRESHOLD_FRAC  = 0.25
N_ITER          = 4
MORPH_CLOSE     = 3
MORPH_OPEN      = 3
GAUSS_BLUR      = 3
NOZZLE_CENTER   = 552
N_BIGGEST       = 2
MIN_CONTOUR_PX  = 20
SAVE_DEBUG_BINARY = True
SAVE_OVERLAY      = True
OVERLAY_COUNT     = 5
SAVE_POSITION_CBAR_MAPS = True   # per-position c̄ map with contours

# c̄ threshold levels — one contour line per value
CBAR_LEVELS     = [0.4, 0.5, 0.6, 0.7]
CLOSE_THRESHOLD = 5
MARK_FLAME_HEIGHT = True
POS1_OFFSET_MM    = 20   # nozzle exit → bottom of pos1 [mm]; 0 disables pos0

# ╔════════════════════════════════════════════════════════════════════╗
# ║                  END OF CONFIG                                     ║
# ╚════════════════════════════════════════════════════════════════════╝

# Machine-specific overrides (paths, workers, …) live in config_local.py,
# which git ignores — so `git pull` never overwrites your settings.
# Copy config_local.example.py → config_local.py and edit it there.
try:
    from config_local import *  # noqa: F401,F403
    _CONFIG_LOCAL = os.path.join(os.path.dirname(os.path.abspath(__file__)), "config_local.py")
except ImportError:
    _CONFIG_LOCAL = None

SCRIPT_VERSION = "v8.3"

# Settings baked into per-position products (logged, never auto-invalidated)
PROCESSING_KEYS = ["SPATIAL_RES", "CROP_BOTTOM_MM", "CROP_TOP_MM", "NOISE_FLOOR",
                   "WINDOW_SIZE", "THRESHOLD_FRAC", "N_ITER", "MORPH_CLOSE",
                   "MORPH_OPEN", "GAUSS_BLUR", "MAX_FRAMES", "MISSING_BG_ACTION"]
INSTANT_KEYS    = ["N_INSTANT", "NOZZLE_CENTER", "N_BIGGEST", "MIN_CONTOUR_PX"]
# Settings applied at flame level (cheap, recomputed when changed)
ANALYSIS_KEYS   = ["DO_CBAR", "DO_INSTANT", "CBAR_LEVELS", "CLOSE_THRESHOLD",
                   "NOZZLE_CENTER", "POS1_OFFSET_MM", "SPATIAL_RES",
                   "NOZZLE_DIAM_FALLBACK_MM", "ALLOW_POSITION_GAPS",
                   "N_INSTANT_COMPOSITES", "STITCH_BAND_PX", "MARK_FLAME_HEIGHT",
                   "INSTANT_AREA", "GAP_CLOSE_PX", "MIN_POCKET_AREA_PX", "FRONT_SMOOTH_PX"]

AVG_DIR_NAME  = "average composites"
INST_DIR_NAME = "instantaneous"            # raw (bg-subtracted) + front
INST_BIN_DIR_NAME = "instantaneous_binarized"  # binary + front


def current_settings(keys):
    g = globals()
    return OrderedDict((k, g[k]) for k in keys)


def product_settings(product):
    keys = PROCESSING_KEYS + (INSTANT_KEYS if product == "instant" else [])
    return current_settings(keys)


# ══════════════════════════════════════════════════════════════════════
#  Logging
# ══════════════════════════════════════════════════════════════════════
class _Tee:
    def __init__(self, path):
        self.f = open(path, "a", encoding="utf-8")
        self.out = sys.__stdout__

    def write(self, s):
        self.out.write(s)
        self.f.write(s)

    def flush(self):
        self.out.flush()
        self.f.flush()


WARNINGS = []   # (flame, message) collected in the main process


def _ts():
    return datetime.datetime.now().strftime("%H:%M:%S")


def log(msg=""):
    print(msg, flush=True)


def warn(flame, msg):
    WARNINGS.append((flame, msg))
    print(f"[{_ts()}]   ⚠ WARNING [{flame}] {msg}", flush=True)


def fmt_dur(sec):
    sec = int(max(sec, 0))
    return f"{sec // 3600}:{(sec % 3600) // 60:02d}:{sec % 60:02d}"


# ══════════════════════════════════════════════════════════════════════
#  Image processing (identical maths to v7)
# ══════════════════════════════════════════════════════════════════════
def crop_px():
    return (int(round(CROP_BOTTOM_MM / SPATIAL_RES)),
            int(round(CROP_TOP_MM / SPATIAL_RES)))


def to_uint8(img):
    img = img.astype(np.float64); img -= img.min(); mx = img.max()
    if mx > 0: img = img / mx * 255.0
    return img.clip(0, 255).astype(np.uint8)


def wabel_local_threshold(img, ws, frac, n):
    work = img.astype(np.float64).copy()
    for _ in range(n):
        lm = maximum_filter(work, size=ws); work[work < frac * lm] = 0.0
    return work > 0


def crop_frame(frame):
    cb, ct = crop_px()
    h = frame.shape[0]; t = ct; b = h - cb
    return frame[t:b] if b > t else frame


def binary_u8(img):
    """v7 make_binary / detect_contours preprocessing → uint8 0/255."""
    work = img.astype(np.float64)
    fl = NOISE_FLOOR / 100.0 * work.max(); work[work < fl] = 0.0
    bi = wabel_local_threshold(work.clip(0, 255).astype(np.uint8),
                               WINDOW_SIZE, THRESHOLD_FRAC, N_ITER)
    u8 = bi.astype(np.uint8) * 255
    if MORPH_CLOSE > 0:
        k = cv2.getStructuringElement(cv2.MORPH_ELLIPSE, (MORPH_CLOSE, MORPH_CLOSE))
        u8 = cv2.morphologyEx(u8, cv2.MORPH_CLOSE, k)
    if MORPH_OPEN > 0:
        k = cv2.getStructuringElement(cv2.MORPH_ELLIPSE, (MORPH_OPEN, MORPH_OPEN))
        u8 = cv2.morphologyEx(u8, cv2.MORPH_OPEN, k)
    if GAUSS_BLUR > 0:
        u8 = cv2.GaussianBlur(u8, (GAUSS_BLUR, GAUSS_BLUR), 0)
        _, u8 = cv2.threshold(u8, 127, 255, cv2.THRESH_BINARY)
    return u8


def make_binary(img):
    return binary_u8(img) > 0


def _split_contour_both(cnt, nc):
    pts = cnt.squeeze()
    if pts.ndim != 2 or len(pts) < 6: return pts, np.empty((0, 2))
    r = pts[:, 1]; ni = int(np.argmin(r)); si = int(np.argmax(r))
    if ni > si: ni, si = si, ni
    a = pts[ni:si + 1]; b = np.vstack([pts[si:], pts[:ni + 1]])
    if len(a) < 2 or len(b) < 2: return pts, np.empty((0, 2))
    da = abs(a[:, 0].mean() - nc); db = abs(b[:, 0].mean() - nc)
    return (a, b) if da < db else (b, a)


def _split_contour_inner(cnt, nc):
    i, _ = _split_contour_both(cnt, nc); return i


def front_from_u8(u8):
    """v7 detect_contours front (red overlay line) from the binary image."""
    cnts, hier = cv2.findContours(u8.copy(), cv2.RETR_CCOMP, cv2.CHAIN_APPROX_NONE)
    F = np.zeros_like(u8, dtype=np.uint8)
    if hier is None: return F > 0
    ext = []; hol = []
    for i, c in enumerate(cnts):
        if hier[0][i][3] == -1: ext.append((c, cv2.contourArea(c)))
        else: hol.append(c)
    ext.sort(key=lambda x: x[1], reverse=True)
    for c, a in ext[:N_BIGGEST]:
        if len(c) < MIN_CONTOUR_PX: continue
        ip = _split_contour_inner(c, NOZZLE_CENTER)
        if len(ip) >= 2:
            cv2.polylines(F, [ip.reshape(-1, 1, 2).astype(np.int32)], False, 255, 1)
    for c, a in ext[N_BIGGEST:]:
        if len(c) >= MIN_CONTOUR_PX: cv2.drawContours(F, [c], -1, 255, 1)
    for c in hol:
        if len(c) >= MIN_CONTOUR_PX: cv2.drawContours(F, [c], -1, 255, 1)
    return F > 0


def save_overlay(u8, front, path):
    rgb = cv2.cvtColor(u8, cv2.COLOR_GRAY2BGR)
    rgb[front, 2] = 255; rgb[front, 1] = 0; rgb[front, 0] = 0
    cv2.imwrite(path, rgb)


# ══════════════════════════════════════════════════════════════════════
#  c̄ contours (identical results to v7, vectorised row search + status)
# ══════════════════════════════════════════════════════════════════════
def scan_cbar_contour_status(cbar, nc, level):
    """v7 ray-scan.  Extra: returns info dict with termination status:
       'closed'      – left/right met (flame tip found)       → OK
       'reached_top' – top row reached without closing         → tip outside FOV
       'no_crossing' – row with no crossing on either side     → open tip
       'empty'       – no contour point at all"""
    h, w = cbar.shape
    left_pts, right_pts = [], []
    status = "reached_top"; single = 0
    for y in range(h - 1, -1, -1):
        lseg = cbar[y, nc::-1] >= level
        rseg = cbar[y, nc:] >= level
        lx = nc - int(np.argmax(lseg)) if lseg.any() else None
        rx = nc + int(np.argmax(rseg)) if rseg.any() else None
        if lx is None and rx is None:
            status = "no_crossing"; break
        if (lx is None) != (rx is None):
            single += 1
        if lx is not None: left_pts.append([lx, y])
        if rx is not None: right_pts.append([rx, y])
        if lx is not None and rx is not None and abs(rx - lx) < CLOSE_THRESHOLD:
            status = "closed"; break
    if not left_pts and not right_pts:
        status = "empty"
    left_pts.reverse(); right_pts.reverse()
    left = np.array(left_pts, dtype=np.int32) if left_pts else np.empty((0, 2), dtype=np.int32)
    right = np.array(right_pts, dtype=np.int32) if right_pts else np.empty((0, 2), dtype=np.int32)
    return left, right, {"status": status, "single_side_rows": single,
                         "rows": max(len(left_pts), len(right_pts))}


def scan_cbar_contour(cbar, nc, level):
    l, r, _ = scan_cbar_contour_status(cbar, nc, level)
    return l, r


def contour_surface_area(left, right, nc, res):
    A = 0.0
    for pts in [left, right]:
        if len(pts) < 2:
            continue
        for j in range(len(pts) - 1):
            x0, y0 = float(pts[j, 0]), float(pts[j, 1])
            x1, y1 = float(pts[j + 1, 0]), float(pts[j + 1, 1])
            r = abs(((x0 + x1) / 2.0) - nc) * res
            ds = np.sqrt((x1 - x0) ** 2 + (y1 - y0) ** 2) * res
            A += np.pi * r * ds
    return A


def contour_flame_height(contours_by_level, img_h, res):
    tips = {}
    for level, (left, right) in contours_by_level.items():
        min_row = img_h - 1
        tip_col = 0
        for pts in [left, right]:
            if len(pts) > 0:
                idx = int(np.argmin(pts[:, 1]))
                if pts[idx, 1] < min_row:
                    min_row = int(pts[idx, 1])
                    tip_col = int(pts[idx, 0])
        height_mm = (img_h - 1 - min_row) * res
        tips[level] = (height_mm, min_row, tip_col)
    return tips


def extract_nozzle_diameter(flame):
    m = re.search(r"(?:^|[\s_\-])d(\d+)(?=$|[\s_\-])", flame, re.IGNORECASE)
    return int(m.group(1)) if m else None


def make_nozzle_extension(width, offset_mm, res):
    ext_px = int(round(offset_mm / res))
    if ext_px <= 0:
        return np.empty((0, width), dtype=np.float64)
    return np.zeros((ext_px, width), dtype=np.float64)


def extend_contours_to_nozzle(left, right, nc, nozzle_diam_mm, res,
                              composite_h_real, ext_px):
    nozzle_r_px = (nozzle_diam_mm / 2.0) / res
    nozzle_exit_row = composite_h_real + ext_px - 1
    left_lip_x = nc - nozzle_r_px
    right_lip_x = nc + nozzle_r_px

    def _linear_bridge(pts, target_x):
        if len(pts) == 0:
            return pts
        x0, y0 = float(pts[-1, 0]), float(pts[-1, 1])
        x1, y1 = target_x, float(nozzle_exit_row)
        n_rows = int(y1 - y0)
        if n_rows <= 0:
            return pts
        new_pts = []
        for i in range(1, n_rows + 1):
            t = i / n_rows
            new_pts.append([int(round(x0 + t * (x1 - x0))),
                            int(round(y0 + t * (y1 - y0)))])
        return np.vstack([pts, np.array(new_pts, dtype=np.int32)])

    return _linear_bridge(left, left_lip_x), _linear_bridge(right, right_lip_x)


# ══════════════════════════════════════════════════════════════════════
#  Chunked SPE reading
# ══════════════════════════════════════════════════════════════════════
def open_spe_stream(path, max_frames=None, is_bg=False):
    """Returns (info, gen) — gen(chunk) yields (first_index, float32 array
    (k, H, W)).  Leading empty frames are skipped like read_spe_stack()."""
    meta = read_spe_metadata(path)
    w, h, count = meta["width"], meta["height"], meta["count"]
    dtype = np.dtype(meta["dtype"])
    fb = w * h * dtype.itemsize
    notes = []
    if meta.get("count_xml_declared") is not None:
        notes.append(f"XML footer declares {meta['count_xml_declared']} frames but "
                     f"the file holds {count} → using {count}")
    lo = meta.get("leftover_bytes")
    if lo and lo > fb:
        notes.append(f"{lo:,} unexplained bytes between data and XML footer "
                     f"(more than one frame) — file may be damaged")
    if count <= 0:
        raise ValueError("file contains 0 frames")
    start = _first_real_frame_index(path, meta["data_start_offset"], fb, count)
    if start >= count:
        if is_bg:
            start = 0
            notes.append("all bg frames are below the 'empty frame' threshold "
                         "(dark background) → using all frames")
        else:
            raise ValueError(f"no frame of the {count}-frame stack has signal — "
                             f"acquisition empty or corrupt")
    if start > 0:
        notes.append(f"skipped {start} leading empty frames")
    n = count - start
    if max_frames is not None:
        n = min(n, max_frames)
    off = meta["data_start_offset"] + start * fb

    def gen(chunk):
        with open(path, "rb") as f:
            f.seek(off)
            done = 0
            while done < n:
                k = min(chunk, n - done)
                buf = f.read(fb * k)
                if len(buf) != fb * k:
                    raise IOError(f"file truncated at frame {start + done} "
                                  f"(expected {fb * k:,} bytes, got {len(buf):,})")
                arr = np.frombuffer(buf, dtype=dtype).reshape(k, h, w).astype(np.float32)
                yield done, arr
                done += k

    info = {"width": w, "height": h, "count_in_file": count, "skipped_leading": start,
            "n_frames": n, "source": meta["source"], "notes": notes}
    return info, gen


def load_bg_mean(path):
    info, gen = open_spe_stream(path, is_bg=True)
    acc = None
    for _, arr in gen(CHUNK_FRAMES):
        s = arr.sum(axis=0, dtype=np.float64)
        acc = s if acc is None else acc + s
    return (acc / info["n_frames"]).astype(np.float32), info


# ══════════════════════════════════════════════════════════════════════
#  Small IO helpers
# ══════════════════════════════════════════════════════════════════════
def atomic_write_json(path, obj):
    tmp = path + ".tmp"
    with open(tmp, "w", encoding="utf-8") as f:
        json.dump(obj, f, indent=2, default=str)
    os.replace(tmp, path)


def atomic_save_npy(path, arr):
    tmp = path + ".tmp"
    with open(tmp, "wb") as f:
        np.save(f, arr)
    os.replace(tmp, path)


def load_json(path):
    try:
        with open(path, "r", encoding="utf-8") as f:
            return json.load(f)
    except Exception:
        return None


def safe_name(s):
    return re.sub(r'[\\/:*?"<>|]+', "_", s).strip(" ._") or "unnamed"


def pos_list_str(ps):
    ps = sorted(ps)
    return "pos" + ", ".join(str(p) for p in ps)


def write_position_settings_txt(path, man):
    L = [f"# Settings for {man['flame']} pos{man['pos']}",
         f"# data file(s): {', '.join(n for n, _ in man['sources']['data'])}",
         f"# bg file     : {man['sources']['bg'][0] if man['sources']['bg'] else 'NONE (no bg subtraction)'}",
         ""]
    for prod, p in sorted(man["products"].items()):
        L.append(f"[{prod}]  finished {p['finished']}  frames used = {p['n_frames']}"
                 f"  (skipped leading empty = {p.get('skipped_leading', 0)})")
        for k, v in p["settings"].items():
            L.append(f"{k} = {v!r}")
        L.append("")
    with open(path, "w", encoding="utf-8") as f:
        f.write("\n".join(L))


# ══════════════════════════════════════════════════════════════════════
#  Worker: one position
# ══════════════════════════════════════════════════════════════════════
def _init_worker():
    try:
        cv2.setNumThreads(1)
    except Exception:
        pass


def process_position(job):
    t0 = time.time()
    res = {"flame": job["flame"], "pos": job["pos"], "ok": False, "warnings": [],
           "error": None, "needs": job["needs"], "n_frames": 0, "seconds": 0.0}
    W = res["warnings"]
    try:
        pd = job["pos_dir"]; os.makedirs(pd, exist_ok=True)
        need_cbar = "cbar" in job["needs"]
        need_inst = "instant" in job["needs"]

        bg = None
        if job["bg_path"]:
            bg, bg_info = load_bg_mean(job["bg_path"])
            atomic_save_npy(os.path.join(pd, "bg_mean.npy"), bg)
            for n in bg_info["notes"]:
                W.append(f"bg '{os.path.basename(job['bg_path'])}': {n}")
        else:
            W.append("NO BACKGROUND FILE → frames processed WITHOUT bg subtraction")

        if need_cbar:
            limit = MAX_FRAMES
        else:
            limit = N_INSTANT if MAX_FRAMES is None else min(N_INSTANT, MAX_FRAMES)

        bsum = None; n_used = 0; n_empty = 0; n_full = 0; skipped = 0
        fronts, bins, fidx = [], [], []
        for spe in job["spe_paths"]:
            remaining = None if limit is None else limit - n_used
            if remaining is not None and remaining <= 0:
                break
            info, gen = open_spe_stream(spe, remaining)
            name = os.path.basename(spe)
            for n in info["notes"]:
                W.append(f"'{name}': {n}")
            skipped += info["skipped_leading"]
            bg_use = None
            if bg is not None:
                fh, fw = info["height"], info["width"]
                if bg.shape == (fh, fw):
                    bg_use = bg
                elif bg.shape[0] >= fh and bg.shape[1] >= fw:
                    W.append(f"bg shape {bg.shape} ≠ frame shape {(fh, fw)} → bg cropped")
                    bg_use = bg[:fh, :fw]
                else:
                    raise ValueError(f"bg shape {bg.shape} smaller than frame "
                                     f"shape {(fh, fw)}")
            for i0, arr in gen(CHUNK_FRAMES):
                if bg_use is not None:
                    arr -= bg_use[np.newaxis, :, :]
                    np.maximum(arr, 0, out=arr)
                for j in range(arr.shape[0]):
                    g = n_used + i0 + j
                    img = to_uint8(crop_frame(arr[j]))
                    u8 = binary_u8(img); bi = u8 > 0
                    if need_cbar:
                        if bsum is None:
                            bsum = np.zeros(bi.shape, dtype=np.float64)
                        bsum += bi
                    fr = bi.mean()
                    if fr < 0.001: n_empty += 1
                    elif fr > 0.999: n_full += 1
                    if need_inst and len(fronts) < N_INSTANT:
                        fronts.append(front_from_u8(u8)); bins.append(bi); fidx.append(g)
                    if g < OVERLAY_COUNT:
                        if SAVE_OVERLAY:
                            save_overlay(img, front_from_u8(u8),
                                         os.path.join(pd, f"overlay_{g:04d}.png"))
                        if SAVE_DEBUG_BINARY:
                            cv2.imwrite(os.path.join(pd, f"binary_{g:04d}.png"), u8)
            n_used += info["n_frames"]

        if n_used == 0:
            raise ValueError("no frames read")
        res["n_frames"] = n_used
        if need_cbar and MAX_FRAMES is None and n_used != EXPECTED_FRAMES:
            W.append(f"{n_used} frames used for c̄ (expected {EXPECTED_FRAMES})")
        if n_empty:
            W.append(f"{n_empty}/{n_used} frames have an (almost) EMPTY binary "
                     f"(<0.1 % burnt) — check signal / threshold")
        if n_full:
            W.append(f"{n_full}/{n_used} frames are >99.9 % 'burnt' — check bg / threshold")

        now = datetime.datetime.now().isoformat(timespec="seconds")
        products = dict(job["keep_products"])
        if need_cbar:
            atomic_save_npy(os.path.join(pd, "cbar.npy"), bsum / n_used)
            products["cbar"] = {"settings": product_settings("cbar"), "finished": now,
                                "n_frames": n_used, "skipped_leading": skipped}
        if need_inst:
            tmp = os.path.join(pd, "instant_fronts.npz.tmp")
            with open(tmp, "wb") as f:
                np.savez_compressed(f, fronts=np.array(fronts, dtype=bool),
                                    binaries=np.array(bins, dtype=bool),
                                    frame_index=np.array(fidx, dtype=np.int32))
            os.replace(tmp, os.path.join(pd, "instant_fronts.npz"))
            products["instant"] = {"settings": product_settings("instant"), "finished": now,
                                   "n_frames": len(fronts), "n_requested": N_INSTANT,
                                   "skipped_leading": skipped}
        man = {"flame": job["flame"], "pos": job["pos"], "sources": job["sources"],
               "products": products, "script_version": SCRIPT_VERSION}
        write_position_settings_txt(os.path.join(pd, f"settings_pos{job['pos']}.txt"), man)
        atomic_write_json(os.path.join(pd, "done.json"), man)   # written LAST
        res["ok"] = True
    except Exception as e:
        res["error"] = f"{type(e).__name__}: {e}\n{traceback.format_exc()}"
    res["seconds"] = time.time() - t0
    return res


# ══════════════════════════════════════════════════════════════════════
#  Discovery & resume scan
# ══════════════════════════════════════════════════════════════════════
NAME_RE = re.compile(r"^(?P<flame>.*?)[\s_\-]*pos(?P<pos>\d+)(?P<rest>.*)$", re.IGNORECASE)
BG_RE = re.compile(r"(?:^|[\s_\-.])bg(?:[\s_\-.]|$)", re.IGNORECASE)


def discover(input_dir):
    raw, ignored = {}, []
    for root, dirs, files in os.walk(input_dir):
        dirs.sort()
        for fn in sorted(files):
            if not fn.lower().endswith(".spe"):
                continue
            m = NAME_RE.match(fn[:-4])
            p = os.path.join(root, fn)
            if not m or not m.group("flame").strip(" _-."):
                ignored.append(p); continue
            flame = m.group("flame").strip(" _-.")
            pos = int(m.group("pos"))
            kind = "bg" if BG_RE.search(m.group("rest")) else "data"
            raw.setdefault(flame, {}).setdefault(pos, {"data": [], "bg": []})[kind].append(p)
    return raw, ignored


def resolve_flames(raw):
    """Pick bg per position, collect discovery warnings."""
    flames = OrderedDict()
    for flame in sorted(raw):
        fw = []
        positions = OrderedDict()
        for pos in sorted(raw[flame]):
            ent = raw[flame][pos]
            if not ent["data"]:
                fw.append(f"pos{pos}: only a bg file, no data file → ignored")
                continue
            data_files = sorted(ent["data"])
            if len(data_files) > 1:
                pol = str(DUPLICATE_POLICY).lower()
                if pol in ("newest", "oldest"):
                    keep = data_files[-1] if pol == "newest" else data_files[0]
                    fw.append(f"DUPLICATE pos{pos}: {len(data_files)} recordings → using "
                              f"{pol.upper()} only: {_rel(keep)}  (ignored: "
                              + "; ".join(_rel(p) for p in data_files if p != keep) + ")")
                    data_files = [keep]
                else:
                    fw.append(f"DUPLICATE pos{pos}: {len(data_files)} data files → c̄ is "
                              f"AVERAGED over all of them: "
                              + "; ".join(_rel(p) for p in data_files))
            bg = None
            if ent["bg"]:
                ddir = os.path.dirname(data_files[-1] if str(DUPLICATE_POLICY).lower() == "newest"
                                       else data_files[0])
                same = [b for b in ent["bg"] if os.path.dirname(b) == ddir]
                cands = same or ent["bg"]
                bg = sorted(cands)[0]
                if len(cands) > 1:
                    fw.append(f"pos{pos}: {len(cands)} bg files, using "
                              f"'{os.path.basename(bg)}'")
            else:
                fw.append(f"pos{pos}: NO BACKGROUND FILE found "
                          f"(MISSING_BG_ACTION='{MISSING_BG_ACTION}')")
            positions[pos] = {"data": data_files, "bg": bg}
        if not positions:
            continue
        folders = OrderedDict()
        for pos, e in positions.items():
            for p in e["data"]:
                folders.setdefault(_rel(os.path.dirname(p)) or ".", set()).add(pos)
        if len(folders) > 1:
            fw.append("FLAME FOUND IN SEVERAL FOLDERS (same name on different days?) → "
                      "combined into ONE flame: "
                      + "; ".join(f"{d}: {pos_list_str(ps_)}" for d, ps_ in folders.items()))
        ps = list(positions)
        gaps = [p for p in range(1, max(ps) + 1) if p not in positions]
        if gaps:
            fw.append(f"MISSING POSITIONS {gaps} — composite would stack non-adjacent "
                      f"positions (ALLOW_POSITION_GAPS={ALLOW_POSITION_GAPS})")
        d = extract_nozzle_diameter(flame)
        if d is None:
            if NOZZLE_DIAM_FALLBACK_MM is not None:
                d = NOZZLE_DIAM_FALLBACK_MM
                fw.append(f"nozzle diameter not found in name → fallback {d} mm")
            else:
                fw.append("nozzle diameter NOT found in name and no fallback → "
                          "pos0 extension disabled, A/h exclude the nozzle gap")
        flames[flame] = {"positions": positions, "gaps": gaps, "D": d,
                         "disc_warnings": fw, "folders": list(folders)}
    return flames


def _rel(path):
    """Path relative to INPUT_DIR (shows the day-folder) for messages."""
    try:
        return os.path.relpath(path, INPUT_DIR)
    except ValueError:
        return path


def source_sig(ent):
    return {"data": [[os.path.basename(p), os.path.getsize(p)] for p in ent["data"]],
            "bg": ([os.path.basename(ent["bg"]), os.path.getsize(ent["bg"])]
                   if ent["bg"] else None)}


def settings_diff(old, new):
    return [f"{k}: {old.get(k)!r} → current {new.get(k)!r}"
            for k in new if old.get(k) != new.get(k)]


def plan_position(flame_dir, pos, ent):
    """Return (needs, keep_products, notes)."""
    pd = os.path.join(flame_dir, f"pos{pos}")
    man = load_json(os.path.join(pd, "done.json"))
    sig = json.loads(json.dumps(source_sig(ent)))
    notes = []
    valid = {}
    if man is not None:
        if man.get("sources") == sig:
            valid = man.get("products", {})
        else:
            notes.append("source files changed since last run → reprocess")
    needs, keep = [], {}
    if DO_CBAR:
        p = valid.get("cbar")
        if p and os.path.isfile(os.path.join(pd, "cbar.npy")):
            keep["cbar"] = p
            d = settings_diff(p["settings"], product_settings("cbar"))
            if d:
                notes.append("c̄ done with different settings (NOT reprocessed): " + "; ".join(d))
        else:
            needs.append("cbar")
    if DO_INSTANT:
        p = valid.get("instant")
        if (p and os.path.isfile(os.path.join(pd, "instant_fronts.npz"))
                and p.get("n_requested", 0) >= N_INSTANT):
            keep["instant"] = p
        else:
            needs.append("instant")
    # keep products not requested this run (e.g. instant while DO_INSTANT off)
    for k, v in valid.items():
        if k not in keep and k not in needs:
            if k == "cbar" and os.path.isfile(os.path.join(pd, "cbar.npy")):
                keep[k] = v
            if k == "instant" and os.path.isfile(os.path.join(pd, "instant_fronts.npz")):
                keep[k] = v
    return needs, keep, notes


# ══════════════════════════════════════════════════════════════════════
#  Visualisation
# ══════════════════════════════════════════════════════════════════════
LEVEL_COLORS = ['black', 'limegreen', 'dodgerblue', 'orange', 'magenta',
                'cyan', 'red', 'yellow']
MARKER_COLORS = ['dimgray', 'forestgreen', 'royalblue', 'darkorange', 'darkmagenta',
                 'darkcyan', 'darkred', 'gold']


def level_color(level):
    i = sorted(CBAR_LEVELS).index(level)
    return LEVEL_COLORS[i % len(LEVEL_COLORS)], MARKER_COLORS[i % len(MARKER_COLORS)]


def draw_cbar(cbar, contours_by_level, tips, path, pos_label=None):
    fig, ax = plt.subplots(1, 1, figsize=(10, cbar.shape[0] / cbar.shape[1] * 10), dpi=150)
    im = ax.imshow(cbar, cmap='RdBu_r', vmin=0, vmax=1, aspect='equal',
                   interpolation='bilinear')
    ax.axvline(x=NOZZLE_CENTER, color='goldenrod', linestyle='--', linewidth=1, alpha=0.8)
    legend_items = []
    for level in sorted(contours_by_level):
        left, right = contours_by_level[level]
        color, mcolor = level_color(level)
        af = contour_surface_area(left, right, NOZZLE_CENTER, SPATIAL_RES)
        for pts in [left, right]:
            if len(pts) >= 2:
                ax.plot(pts[:, 0], pts[:, 1], '-', color=color, linewidth=1.5, alpha=0.9)
        legend_items.append(Line2D([0], [0], color=color, linewidth=1.5,
                                   label=f'c̄={level} (A_F={af:.0f} mm²)'))
        if MARK_FLAME_HEIGHT and level in tips:
            hm, tr, tc = tips[level]
            if tr < cbar.shape[0] - 1:
                ax.plot(tc, tr, 'o', color=mcolor, markersize=5, zorder=5)
    if legend_items:
        ax.legend(handles=legend_items, loc='upper right', fontsize=8, framealpha=0.8)
    cb = fig.colorbar(im, ax=ax, fraction=0.03, pad=0.02); cb.set_label('c̄', fontsize=12)
    if pos_label:
        ax.text(cbar.shape[1] - 5, cbar.shape[0] - 10, pos_label,
                ha='right', va='bottom', fontsize=8, color='orange')
    ax.set_xlim(0, cbar.shape[1]); ax.set_ylim(cbar.shape[0], 0); ax.axis('off')
    fig.tight_layout(); fig.savefig(path, bbox_inches='tight', dpi=150); plt.close(fig)


def plot_avg_composite(C, levels, path, title):
    comp = C["composite"]
    fig, ax = plt.subplots(1, 1, figsize=(10, comp.shape[0] / comp.shape[1] * 10), dpi=150)
    im = ax.imshow(comp, cmap='RdBu_r', vmin=0, vmax=1, aspect='equal',
                   interpolation='bilinear')
    ax.axvline(x=NOZZLE_CENTER, color='goldenrod', linestyle='--', linewidth=1, alpha=0.8)
    for y_off, pos in C["pos_boundaries"]:
        if y_off > 0:
            ax.axhline(y=y_off, color='orange', linestyle=':', linewidth=0.5, alpha=0.6)
        ax.text(comp.shape[1] - 5, y_off + 15, f'pos{pos}', ha='right', va='top',
                fontsize=7, color='orange', alpha=0.8)
    if C["ext_px"] > 0:
        yb = C["h_real"]
        ax.axhline(y=yb, color='orange', linestyle=':', linewidth=0.5, alpha=0.6)
        ax.text(comp.shape[1] - 5, yb + 15, 'pos0 (ext.)', ha='right', va='top',
                fontsize=7, color='orange', alpha=0.8)
        r_px = (C["nozzle_diam"] / 2.0) / SPATIAL_RES
        lip_y = comp.shape[0] - 1
        for xx in (NOZZLE_CENTER - r_px, NOZZLE_CENTER + r_px):
            ax.plot(xx, lip_y, 's', color='gray', markersize=5, zorder=5)
    legend_items = []
    for level in sorted(levels):
        left, right = C["contours"][level]
        color, mcolor = level_color(level)
        for pts in [left, right]:
            if len(pts) >= 2:
                ax.plot(pts[:, 0], pts[:, 1], '-', color=color, linewidth=1.5, alpha=0.9)
        legend_items.append(Line2D([0], [0], color=color, linewidth=1.5,
                                   label=f'c̄={level} (A={C["areas"][level][2]:.0f} mm², '
                                         f'h={C["tips"][level][0]:.1f} mm)'))
        if MARK_FLAME_HEIGHT:
            hm, tr, tc = C["tips"][level]
            if tr < comp.shape[0] - 1:
                ax.plot(tc, tr, 'o', color=mcolor, markersize=6, zorder=5)
    if legend_items:
        ax.legend(handles=legend_items, loc='upper right', fontsize=9, framealpha=0.8)
    ax.set_title(title, fontsize=9)
    cb = fig.colorbar(im, ax=ax, fraction=0.03, pad=0.02); cb.set_label('c̄', fontsize=12)
    ax.set_xlim(0, comp.shape[1]); ax.set_ylim(comp.shape[0], 0); ax.axis('off')
    fig.tight_layout(); fig.savefig(path, bbox_inches='tight', dpi=150); plt.close(fig)


# ══════════════════════════════════════════════════════════════════════
#  Flame finalisation — average (c̄) part
# ══════════════════════════════════════════════════════════════════════
def build_cbar_composite(pcb, nozzle_diam):
    sorted_pos = sorted(pcb)
    strips, pos_boundaries, y = [], [], 0
    for pos in reversed(sorted_pos):
        strips.append(pcb[pos]); pos_boundaries.append((y, pos)); y += pcb[pos].shape[0]
    composite_real = np.vstack(strips)
    h_real, w = composite_real.shape

    real, scaninfo = {}, {}
    for level in CBAR_LEVELS:
        l, r, inf = scan_cbar_contour_status(composite_real, NOZZLE_CENTER, level)
        real[level] = (l, r); scaninfo[level] = inf

    ext_px = 0
    composite = composite_real
    if POS1_OFFSET_MM > 0 and nozzle_diam is not None:
        ext = make_nozzle_extension(w, POS1_OFFSET_MM, SPATIAL_RES)
        ext_px = ext.shape[0]
        if ext_px > 0:
            composite = np.vstack([composite_real, ext])
    if ext_px > 0:
        full = {lv: extend_contours_to_nozzle(real[lv][0], real[lv][1], NOZZLE_CENTER,
                                              nozzle_diam, SPATIAL_RES, h_real, ext_px)
                for lv in CBAR_LEVELS}
    else:
        full = real
    tips = contour_flame_height(full, composite.shape[0], SPATIAL_RES)
    height_offset = 0.0
    if ext_px == 0 and POS1_OFFSET_MM > 0:
        # no pos0 rows → heights measured from pos1 bottom; add the gap
        height_offset = float(POS1_OFFSET_MM)
        tips = {lv: (t[0] + height_offset, t[1], t[2]) for lv, t in tips.items()}

    areas = {}
    for lv in CBAR_LEVELS:
        a_real = contour_surface_area(real[lv][0], real[lv][1], NOZZLE_CENTER, SPATIAL_RES)
        a_tot = contour_surface_area(full[lv][0], full[lv][1], NOZZLE_CENTER, SPATIAL_RES)
        areas[lv] = (a_real, a_tot - a_real, a_tot)
    return {"composite": composite, "h_real": h_real, "ext_px": ext_px,
            "pos_boundaries": pos_boundaries, "contours": full, "contours_real": real,
            "scaninfo": scaninfo, "tips": tips, "areas": areas,
            "nozzle_diam": nozzle_diam, "height_offset": height_offset}


SCAN_MSG = {
    "reached_top": "c-line did NOT close before the top of the highest position → "
                   "flame tip outside the field of view; A and h are UNDERESTIMATED",
    "no_crossing": "c-line ended without left/right meeting (row without any "
                   "crossing) → open tip, check the c̄ map",
    "empty": "NO contour points found at this level → A = 0",
}


def finalize_cbar(flame, F, fdir, positions):
    pcb = {p: np.load(os.path.join(fdir, f"pos{p}", "cbar.npy")) for p in positions}
    widths = {p: a.shape[1] for p, a in pcb.items()}
    if len(set(widths.values())) > 1:
        raise ValueError(f"positions have different image widths {widths}")
    heights = {p: a.shape[0] for p, a in pcb.items()}
    if len(set(heights.values())) > 1:
        warn(flame, f"positions have different heights {heights}")

    if SAVE_POSITION_CBAR_MAPS:
        for p in positions:
            cs = {lv: scan_cbar_contour(pcb[p], NOZZLE_CENTER, lv) for lv in CBAR_LEVELS}
            tips = contour_flame_height(cs, pcb[p].shape[0], SPATIAL_RES)
            draw_cbar(pcb[p], cs, tips, os.path.join(fdir, f"pos{p}", "cbar_map.png"),
                      pos_label=f"pos{p}")

    C = build_cbar_composite(pcb, F["D"])
    if C["ext_px"] == 0 and POS1_OFFSET_MM > 0:
        warn(flame, "no pos0 extension (nozzle diameter unknown) → A excludes the "
                    f"nozzle gap; h measured from pos1 bottom + {POS1_OFFSET_MM} mm")
    lv_status = {}
    for lv in CBAR_LEVELS:
        inf = C["scaninfo"][lv]
        lv_status[lv] = inf["status"]
        if inf["status"] in SCAN_MSG:
            warn(flame, f"c̄={lv}: {SCAN_MSG[inf['status']]}")
        if inf["rows"] and inf["single_side_rows"] > 0.05 * inf["rows"]:
            warn(flame, f"c̄={lv}: {inf['single_side_rows']} of {inf['rows']} rows have a "
                        f"crossing on ONE side only — asymmetric/clipped flame?")
        a = C["areas"][lv]; h = C["tips"][lv][0]
        log(f"    c̄={lv}: A_real={a[0]:.1f}  A_ext={a[1]:.1f}  A_total={a[2]:.1f} mm²  "
            f"h={h:.1f} mm  [{inf['status']}]")

    ad = os.path.join(fdir, AVG_DIR_NAME); os.makedirs(ad, exist_ok=True)
    base_t = (f"{flame}  (pos{min(positions)} bottom → pos{max(positions)} top, "
              f"NOZZLE_CENTER={NOZZLE_CENTER})")
    plot_avg_composite(C, [], os.path.join(ad, "composite_no_clines.png"), base_t)
    for lv in CBAR_LEVELS:
        plot_avg_composite(C, [lv], os.path.join(ad, f"composite_cbar_{lv}.png"),
                           base_t + f"  c̄={lv}")
    plot_avg_composite(C, list(CBAR_LEVELS), os.path.join(ad, "composite_all_clines.png"),
                       base_t + "  all c-lines")

    # numpy exports
    atomic_save_npy(os.path.join(ad, "cbar_composite.npy"), C["composite"])
    H = C["composite"].shape[0]
    d = {"levels": np.array(CBAR_LEVELS, dtype=float),
         "nozzle_center_px": np.array(NOZZLE_CENTER),
         "spatial_res_mm": np.array(SPATIAL_RES),
         "composite_h_real_px": np.array(C["h_real"]),
         "ext_px": np.array(C["ext_px"]),
         "nozzle_diam_mm": np.array(np.nan if C["nozzle_diam"] is None else C["nozzle_diam"]),
         "height_offset_mm": np.array(C["height_offset"]),
         "positions_top_to_bottom": np.array([p for _, p in C["pos_boundaries"]]),
         "pos_row_offsets_px": np.array([y for y, _ in C["pos_boundaries"]])}
    for lv in CBAR_LEVELS:
        for arm, pts, rpts in (("left", C["contours"][lv][0], C["contours_real"][lv][0]),
                               ("right", C["contours"][lv][1], C["contours_real"][lv][1])):
            d[f"c{lv}_{arm}_px"] = pts
            d[f"c{lv}_{arm}_rayscan_px"] = rpts
            if len(pts):
                r = (pts[:, 0] - NOZZLE_CENTER) * SPATIAL_RES
                z = (H - 1 - pts[:, 1]) * SPATIAL_RES + C["height_offset"]
                d[f"c{lv}_{arm}_mm"] = np.column_stack([r, z])
            else:
                d[f"c{lv}_{arm}_mm"] = np.empty((0, 2))
    tmp = os.path.join(ad, "c_lines.npz.tmp")
    with open(tmp, "wb") as f:
        np.savez(f, **d)
    os.replace(tmp, os.path.join(ad, "c_lines.npz"))

    return {"levels": {str(lv): {"A_real": C["areas"][lv][0], "A_ext": C["areas"][lv][1],
                                 "A_total": C["areas"][lv][2], "h": C["tips"][lv][0],
                                 "tip_row": C["tips"][lv][1], "tip_col": C["tips"][lv][2],
                                 "status": lv_status[lv]} for lv in CBAR_LEVELS},
            "nozzle_diam": C["nozzle_diam"], "ext_px": C["ext_px"]}


# ══════════════════════════════════════════════════════════════════════
#  Flame finalisation — instantaneous part (best-match chaining)
# ══════════════════════════════════════════════════════════════════════
def _spe_real_range(path):
    """(start, n_real, meta) — same leading-empty-frame skip as open_spe_stream."""
    meta = read_spe_metadata(path)
    fb = meta["width"] * meta["height"] * np.dtype(meta["dtype"]).itemsize
    start = _first_real_frame_index(path, meta["data_start_offset"], fb, meta["count"])
    return start, meta["count"] - start, meta, fb


def _position_bg(fdir, p, ent):
    """Mean background of a position (cached as posN/bg_mean.npy)."""
    cache = os.path.join(fdir, f"pos{p}", "bg_mean.npy")
    if os.path.isfile(cache):
        return np.load(cache)
    if ent["bg"] is None:
        return None
    bg, _ = load_bg_mean(ent["bg"])
    atomic_save_npy(cache, bg)
    return bg


def read_raw_instant_frame(ent, g, bg):
    """Frame g (index as stored in instant_fronts.npz, counted over the
    position's data files after leading empty frames) → bg-subtracted,
    clipped at 0, cropped float32 image.  Returns (img, file, frame_in_file)."""
    for path in ent["data"]:
        start, n, meta, fb = _spe_real_range(path)
        if g < n:
            with open(path, "rb") as f:
                f.seek(meta["data_start_offset"] + (start + g) * fb)
                buf = f.read(fb)
            img = np.frombuffer(buf, dtype=np.dtype(meta["dtype"])).reshape(
                meta["height"], meta["width"]).astype(np.float32)
            if bg is not None:
                b = bg if bg.shape == img.shape else bg[:img.shape[0], :img.shape[1]]
                img = np.maximum(img - b, 0)
            return crop_frame(img), os.path.basename(path), int(start + g)
        g -= n
    raise IndexError("instantaneous frame index beyond the data files")


WHITE_RED = LinearSegmentedColormap.from_list(
    "white_red", ["#ffffff", "#ffb3b3", "#ff4d4d", "#ff0000", "#dd0000"])


def inner_front(FR, BI, nc):
    """Keep only the reactant/product front: front pixels bordering the
    unburnt core or enclosed pockets.  Front pixels bordering unburnt
    regions connected to the LEFT/RIGHT image edge (= ambient air outside
    the products) are removed.  Returns (inner_front, removed_front, notes)."""
    notes = []
    B = BI.astype(np.uint8)
    if GAP_CLOSE_PX and GAP_CLOSE_PX > 1:
        k = cv2.getStructuringElement(cv2.MORPH_ELLIPSE, (int(GAP_CLOSE_PX), int(GAP_CLOSE_PX)))
        B = cv2.morphologyEx(B, cv2.MORPH_CLOSE, k)
    U = (B == 0).astype(np.uint8)
    n, lab, stats, _ = cv2.connectedComponentsWithStats(U, connectivity=4)
    side = set(np.unique(lab[:, 0]).tolist()) | set(np.unique(lab[:, -1]).tolist())
    side.discard(0)
    # core = unburnt component at the bottom row nearest the axis
    bot = np.where(U[-1] > 0)[0]
    core = int(lab[-1, bot[np.argmin(np.abs(bot - nc))]]) if bot.size else 0
    if core and core in side:
        notes.append("unburnt core is connected to the outside through a gap → part of "
                     "the inner front may be missing (try a larger GAP_CLOSE_PX)")
        side.discard(core)
    keep_lab = np.zeros(n, dtype=bool)
    for i in range(1, n):
        if i in side:
            continue
        if i != core and stats[i, cv2.CC_STAT_AREA] < MIN_POCKET_AREA_PX:
            continue
        keep_lab[i] = True
    keep = keep_lab[lab]
    near = cv2.dilate(keep.astype(np.uint8), np.ones((5, 5), np.uint8)) > 0
    return FR & near, FR & ~near, notes


def front_revolution_area(front, nc, res, smooth=FRONT_SMOOTH_PX):
    """Surface area of a 1-px front revolved about the axis x = nc, each
    side 180° → A = Σ π·r·ds (same convention as the c-lines).
    The front pixels are traced as contours (a 1-px line is traced once
    along each side, so every piece is counted twice → divided by 2) and
    the coordinates are smoothed along the contour so the pixel
    staircase does not inflate the length."""
    cnts, _ = cv2.findContours(front.astype(np.uint8), cv2.RETR_LIST, cv2.CHAIN_APPROX_NONE)
    A = 0.0
    k = np.ones(int(smooth)) / int(smooth) if smooth and smooth > 1 else None
    for c in cnts:
        pts = c[:, 0, :].astype(float)
        if len(pts) < 2:
            continue
        if k is not None and len(pts) > len(k):
            pad = len(k)
            xs = np.convolve(np.r_[pts[-pad:, 0], pts[:, 0], pts[:pad, 0]], k, "same")[pad:-pad]
            ys = np.convolve(np.r_[pts[-pad:, 1], pts[:, 1], pts[:pad, 1]], k, "same")[pad:-pad]
        else:
            xs, ys = pts[:, 0], pts[:, 1]
        xn, yn = np.roll(xs, -1), np.roll(ys, -1)          # closed contour
        ds = np.hypot(xn - xs, yn - ys) * res
        r = np.abs((xs + xn) / 2.0 - nc) * res
        A += np.sum(np.pi * r * ds)
    return float(A / 2.0)


def _edge_masks(bins, top):
    B = max(1, STITCH_BAND_PX)
    band = bins[:, :B, :] if top else bins[:, -B:, :]
    return band.mean(axis=1) > 0.5          # (n, W) majority over band rows


def _best_chain(costs, avail):
    D = np.where(avail[0], 0.0, np.inf)
    backs = []
    for k, Ck in enumerate(costs):
        tot = D[:, None] + Ck
        b = np.argmin(tot, axis=0)
        Dn = tot[b, np.arange(tot.shape[1])]
        D = np.where(avail[k + 1], Dn, np.inf)
        backs.append(b)
    end = int(np.argmin(D))
    if not np.isfinite(D[end]):
        return None
    chain = [end]
    for k in range(len(costs) - 1, -1, -1):
        chain.append(int(backs[k][chain[-1]]))
    chain.reverse()
    return chain


def finalize_instant(flame, F, fdir, positions):
    data = {}
    for p in positions:
        z = np.load(os.path.join(fdir, f"pos{p}", "instant_fronts.npz"))
        data[p] = (z["fronts"], z["binaries"], z["frame_index"])
    widths = {p: d[0].shape[2] for p, d in data.items()}
    if len(set(widths.values())) > 1:
        raise ValueError(f"instantaneous fronts have different widths {widths}")
    ps = sorted(positions)                       # bottom → top
    # cost[k][a, b]: mismatch (px) between top edge of ps[k] frame a and
    # bottom edge of ps[k+1] frame b
    costs = []
    for k in range(len(ps) - 1):
        L = _edge_masks(data[ps[k]][1], top=True).astype(np.float32)
        U = _edge_masks(data[ps[k + 1]][1], top=False).astype(np.float32)
        costs.append(L.sum(1)[:, None] + U.sum(1)[None, :] - 2.0 * (L @ U.T))
    avail = [np.ones(data[p][0].shape[0], dtype=bool) for p in ps]

    idir = os.path.join(fdir, INST_DIR_NAME); os.makedirs(idir, exist_ok=True)
    bdir = os.path.join(fdir, INST_BIN_DIR_NAME); os.makedirs(bdir, exist_ok=True)
    for stale in (glob.glob(os.path.join(idir, "instant_binary_composite_*.npy"))
                  + ([] if INSTANT_AREA else
                     glob.glob(os.path.join(idir, "instant_front_inner_composite_*.npy"))
                     + glob.glob(os.path.join(bdir, "instant_front_inner_composite_*.npy")))):
        os.remove(stale)        # binary composites now live in instantaneous_binarized/
    bgs = {p: _position_bg(fdir, p, F["positions"][p]) for p in ps}
    out = []
    for ci in range(1, N_INSTANT_COMPOSITES + 1):
        chain = _best_chain(costs, avail)
        if chain is None:
            warn(flame, f"instantaneous: only {ci - 1} composites possible "
                        f"(not enough frames per position)")
            break
        for k, a in enumerate(chain):
            avail[k][a] = False
        jc = [float(costs[k][chain[k], chain[k + 1]]) for k in range(len(costs))]

        fr_strips, bi_strips, raw_strips, srcs, bounds, y = [], [], [], {}, [], 0
        for k in range(len(ps) - 1, -1, -1):        # top position first
            p = ps[k]
            fr = data[p][0][chain[k]].copy()
            # findContours traces regions touching the image border along the
            # border itself — not a flame front.  Remove border pixels so they
            # don't show up as horizontal lines at every position junction.
            fr[0, :] = False; fr[-1, :] = False; fr[:, 0] = False; fr[:, -1] = False
            fr_strips.append(fr); bi_strips.append(data[p][1][chain[k]])
            g = int(data[p][2][chain[k]])
            raw, fname, fidx = read_raw_instant_frame(F["positions"][p], g, bgs[p])
            if raw.shape != fr.shape:
                raise ValueError(f"pos{p}: raw frame shape {raw.shape} ≠ stored front "
                                 f"shape {fr.shape}")
            raw_strips.append(raw)
            srcs[f"pos{p}"] = {"file": fname, "frame_in_file": fidx}
            bounds.append((y, p, g))
            y += data[p][0].shape[1]
        FR = np.vstack(fr_strips); BI = np.vstack(bi_strips); RAW = np.vstack(raw_strips)
        for d_ in (idir, bdir):
            atomic_save_npy(os.path.join(d_, f"instant_front_composite_{ci}.npy"), FR)
        atomic_save_npy(os.path.join(bdir, f"instant_binary_composite_{ci}.npy"), BI)
        atomic_save_npy(os.path.join(idir, f"instant_raw_composite_{ci}.npy"), RAW)
        info = {"flame": flame, "composite": ci,
                "frames_used": {f"pos{p}": int(data[p][2][chain[k]]) for k, p in enumerate(ps)},
                "source_frames": srcs,
                "junction_mismatch_px": {f"pos{ps[k]}|pos{ps[k + 1]}": jc[k]
                                         for k in range(len(jc))},
                "note": "Positions were recorded at different times: this is a "
                        "SYNTHETIC stitched front for visualisation only.",
                "rows_top_to_bottom": [{"row_offset": b[0], "pos": b[1], "frame": b[2]}
                                       for b in bounds]}
        FRrem = np.zeros_like(FR)
        if INSTANT_AREA:
            FRin, FRrem, notes = inner_front(FR, BI, NOZZLE_CENTER)
            for n_ in notes:
                warn(flame, f"instantaneous #{ci}: {n_}")
            info["A_inst_mm2"] = round(front_revolution_area(FRin, NOZZLE_CENTER, SPATIAL_RES), 2)
            info["A_inst_note"] = ("inner front only (outer product/air edge removed), "
                                   "each side revolved 180° about NOZZLE_CENTER="
                                   f"{NOZZLE_CENTER}; covers the measured positions only "
                                   "(compare with A_real, not A_total)")
            info["inner_front_warnings"] = notes
            for d_ in (idir, bdir):
                atomic_save_npy(os.path.join(d_, f"instant_front_inner_composite_{ci}.npy"), FRin)
            FR = FRin
        for d_ in (idir, bdir):
            atomic_write_json(os.path.join(d_, f"instant_composite_{ci}_info.json"), info)

        FRd = cv2.dilate(FR.astype(np.uint8), np.ones((3, 3), np.uint8)) > 0
        FRremd = cv2.dilate(FRrem.astype(np.uint8), np.ones((3, 3), np.uint8)) > 0
        a_txt = (f"   A_inst = {info['A_inst_mm2']:.0f} mm²" if "A_inst_mm2" in info else "")

        def _annotate(ax, txt_color):
            ax.axvline(x=NOZZLE_CENTER, color='goldenrod', linestyle='--', linewidth=1, alpha=0.8)
            for yo, p, fi in bounds:
                if yo > 0:
                    ax.axhline(y=yo, color='orange', linestyle=':', linewidth=0.6, alpha=0.9)
                ax.text(FR.shape[1] - 5, yo + 15, f'pos{p}  frame {fi}', ha='right', va='top',
                        fontsize=7, color='darkorange')
            for k in range(len(jc)):
                yo = [b[0] for b in bounds if b[1] == ps[k + 1]][0] + data[ps[k + 1]][0].shape[1]
                ax.text(5, yo - 5, f'mismatch {jc[k]:.0f} px', ha='left', va='bottom',
                        fontsize=7, color=txt_color)

        # ── raw (bg-subtracted) signal: 0 = white → strong = red, front = blue
        pos_vals = RAW[RAW > 0]
        vmax = float(np.percentile(pos_vals, 99.5)) if pos_vals.size else 1.0
        fig, ax = plt.subplots(1, 1, figsize=(10, FR.shape[0] / FR.shape[1] * 10), dpi=150)
        im = ax.imshow(RAW, cmap=WHITE_RED, vmin=0, vmax=max(vmax, 1e-6),
                       aspect='equal', interpolation='nearest')
        ov = np.zeros(FR.shape + (4,), dtype=float)
        ov[FRremd] = (0.55, 0.55, 0.55, 0.8)       # removed outer edge: grey
        ov[FRd] = (0.0, 0.25, 1.0, 1.0)
        ax.imshow(ov, aspect='equal', interpolation='nearest')
        _annotate(ax, 'black')
        cb = fig.colorbar(im, ax=ax, fraction=0.03, pad=0.02)
        cb.set_label('OH-PLIF signal, bg-subtracted [counts]', fontsize=9)
        ax.set_title(f"{flame} — stitched instantaneous front #{ci}, raw signal{a_txt}\n"
                     f"(synthetic: positions not simultaneous; grey = outer edge, not counted)",
                     fontsize=9)
        ax.axis('off'); fig.tight_layout()
        fig.savefig(os.path.join(idir, f"instant_composite_{ci}.png"),
                    bbox_inches='tight', dpi=150)
        plt.close(fig)

        # ── binarized version (as before)
        rgb = np.zeros(FR.shape + (3,), dtype=np.uint8)
        rgb[BI] = (70, 70, 70)
        rgb[FRremd] = (120, 60, 60)                 # removed outer edge: dim red
        rgb[FRd] = (255, 40, 40)
        fig, ax = plt.subplots(1, 1, figsize=(10, FR.shape[0] / FR.shape[1] * 10), dpi=150)
        ax.imshow(rgb, aspect='equal', interpolation='nearest')
        _annotate(ax, 'cyan')
        ax.set_title(f"{flame} — stitched instantaneous front #{ci}, binarized{a_txt}\n"
                     f"(synthetic: positions not simultaneous; dim red = outer edge, not counted)",
                     fontsize=9)
        ax.axis('off'); fig.tight_layout()
        fig.savefig(os.path.join(bdir, f"instant_composite_{ci}.png"),
                    bbox_inches='tight', dpi=150)
        plt.close(fig)
        out.append(info)
        log(f"    instantaneous #{ci}: frames {info['frames_used']}  "
            f"mismatch px {[round(c) for c in jc]}{a_txt}")
    return out


# ══════════════════════════════════════════════════════════════════════
#  Settings log, Excel
# ══════════════════════════════════════════════════════════════════════
def build_settings_log(flame, fdir, positions, mans):
    L = [f"USER SETTINGS LOG — {flame}",
         f"generated {datetime.datetime.now().isoformat(timespec='seconds')} "
         f"by run_wabel_batch.py {SCRIPT_VERSION}", "",
         "=== Flame-level analysis settings (applied in this run) ==="]
    L += [f"{k} = {v!r}" for k, v in current_settings(ANALYSIS_KEYS).items()]
    L += ["", "=== Per-position processing settings ==="]
    groups = OrderedDict()
    for p in positions:
        key = json.dumps({prod: v["settings"] for prod, v in sorted(mans[p]["products"].items())},
                         sort_keys=True, default=str)
        groups.setdefault(key, []).append(p)
    for key, plist in groups.items():
        hdr = (f"for all positions ({pos_list_str(plist)}):" if len(groups) == 1
               else f"for {pos_list_str(plist)}:")
        L.append(hdr)
        prods = json.loads(key)
        cb = prods.get("cbar")
        for prod, st in prods.items():
            if prod == "instant" and cb is not None and all(
                    st.get(k) == cb.get(k) for k in PROCESSING_KEYS):
                L.append("  [instant]  same image-processing settings as [cbar], plus:")
                L += [f"    {k} = {st.get(k)!r}" for k in INSTANT_KEYS]
                continue
            L.append(f"  [{prod}]")
            L += [f"    {k} = {v!r}" for k, v in st.items()]
        L.append("")
    diffs = []
    for p in positions:
        for prod, v in sorted(mans[p]["products"].items()):
            d = settings_diff(v["settings"], product_settings(prod))
            if d:
                diffs.append(f"pos{p} [{prod}]: " + "; ".join(d))
    L.append("=== Positions processed with settings different from the CURRENT config "
             "(kept, not reprocessed) ===")
    L += diffs or ["none"]
    L += ["", "=== Per-position input ==="]
    for p in positions:
        m = mans[p]
        bg = m["sources"]["bg"][0] if m["sources"]["bg"] else "NONE → no bg subtraction"
        L.append(f"pos{p}: data = {', '.join(n for n, _ in m['sources']['data'])} | bg = {bg}")
        for prod, v in sorted(m["products"].items()):
            L.append(f"        [{prod}] frames used = {v['n_frames']}, skipped leading "
                     f"empty = {v.get('skipped_leading', 0)}, finished {v['finished']}")
    with open(os.path.join(fdir, "user_settings_log.txt"), "w", encoding="utf-8") as f:
        f.write("\n".join(L) + "\n")


def A_hdr(lv): return f"A_(c̄={lv}) [mm²]"
def h_hdr(lv): return f"h_(c̄={lv}) [mm]"


def write_table(path, sheets):
    """sheets: list of (name, header, rows).  xlsx via openpyxl, else CSVs."""
    try:
        from openpyxl import Workbook
        from openpyxl.utils import get_column_letter
        wb = Workbook(); first = True
        for name, header, rows in sheets:
            ws = wb.active if first else wb.create_sheet()
            first = False
            ws.title = name[:31]
            ws.append(header)
            for r in rows:
                ws.append(r)
            for i, hd in enumerate(header, 1):
                ws.column_dimensions[get_column_letter(i)].width = max(12, min(60, len(str(hd)) + 2))
        tmp = path + ".tmp.xlsx"
        wb.save(tmp); os.replace(tmp, path)
        return path
    except ImportError:
        import csv
        base = os.path.splitext(path)[0]
        for name, header, rows in sheets:
            with open(f"{base}_{safe_name(name)}.csv", "w", newline="", encoding="utf-8") as f:
                w = csv.writer(f); w.writerow(header); w.writerows(rows)
        return base + "_*.csv (openpyxl missing)"


def results_row(flame, cb):
    lv_sorted = sorted(CBAR_LEVELS)
    return ([flame] + [round(cb["levels"][str(lv)]["A_total"], 2) for lv in lv_sorted]
            + [round(cb["levels"][str(lv)]["h"], 2) for lv in lv_sorted])


def results_header():
    lv_sorted = sorted(CBAR_LEVELS)
    return ["Condition"] + [A_hdr(lv) for lv in lv_sorted] + [h_hdr(lv) for lv in lv_sorted]


def instant_rows(inst):
    return [[i["composite"], i.get("A_inst_mm2", ""),
             ", ".join(f"{k}:{v}" for k, v in i["frames_used"].items()),
             ", ".join(f"{v:.0f}" for v in i["junction_mismatch_px"].values())]
            for i in (inst or [])]


def write_flame_excel(flame, fdir, positions, mans, cb, inst=None):
    sheets = [("Results", results_header(), [results_row(flame, cb)])]
    if inst and any("A_inst_mm2" in i for i in inst):
        sheets.append(("Instantaneous", ["composite #", "A_inst [mm²] (measured part only)",
                                         "frames used", "junction mismatch [px]"],
                       instant_rows(inst)))
    rows = []
    for lv in sorted(CBAR_LEVELS):
        d = cb["levels"][str(lv)]
        rows.append([lv, round(d["A_real"], 2), round(d["A_ext"], 2), round(d["A_total"], 2),
                     round(d["h"], 2), d["tip_row"], d["tip_col"], d["status"]])
    sheets.append(("Area breakdown", ["c̄ level", "A_real pos1+ [mm²]", "A_ext pos0 [mm²]",
                                      "A_total [mm²]", "h [mm]", "tip row", "tip col",
                                      "contour status"], rows))
    prow = []
    for p in positions:
        m = mans[p]; c = m["products"].get("cbar", {})
        prow.append([p, ", ".join(n for n, _ in m["sources"]["data"]),
                     m["sources"]["bg"][0] if m["sources"]["bg"] else "NONE",
                     c.get("n_frames", ""), c.get("skipped_leading", ""), c.get("finished", "")])
    sheets.append(("Positions", ["pos", "data file(s)", "bg file", "frames (c̄)",
                                 "skipped leading empty", "processed"], prow))
    return write_table(os.path.join(fdir, f"{safe_name(flame)}_results.xlsx"), sheets)


# ══════════════════════════════════════════════════════════════════════
#  Flame finalisation (main process)
# ══════════════════════════════════════════════════════════════════════
def finalize_flame(flame, F):
    fdir = os.path.join(OUTDIR, safe_name(flame))
    positions = sorted(F["positions"])
    mans = {p: load_json(os.path.join(fdir, f"pos{p}", "done.json")) for p in positions}
    missing = [p for p in positions if mans[p] is None]
    if missing:
        raise RuntimeError(f"positions {missing} have no completed results")

    fhash = hashlib.sha1(json.dumps(
        {"analysis": current_settings(ANALYSIS_KEYS), "version": SCRIPT_VERSION,
         "pos": {str(p): mans[p] for p in positions}},
        sort_keys=True, default=str).encode()).hexdigest()
    fdone = load_json(os.path.join(fdir, "flame_done.json"))
    fres_path = os.path.join(fdir, "flame_results.json")
    if (not REBUILD_FINISHED_FLAMES and fdone and fdone.get("hash") == fhash
            and os.path.isfile(fres_path)):
        log(f"[{_ts()}]   {flame}: composites/Excel up to date → skipped")
        return load_json(fres_path)

    log(f"[{_ts()}] ▶ Finalising {flame} ({pos_list_str(positions)}) ...")
    t0 = time.time()
    if F["gaps"] and not ALLOW_POSITION_GAPS:
        raise RuntimeError(f"missing positions {F['gaps']} → composite not built "
                           f"(set ALLOW_POSITION_GAPS=True to force)")
    result = {"flame": flame, "positions": positions, "gaps": F["gaps"],
              "no_bg": [p for p in positions if mans[p]["sources"]["bg"] is None],
              "cbar": None, "instant": None}
    if DO_CBAR:
        if all("cbar" in mans[p]["products"] for p in positions):
            result["cbar"] = finalize_cbar(flame, F, fdir, positions)
        else:
            warn(flame, "c̄ results missing for some positions → no c-line output")
    if DO_INSTANT:
        if all("instant" in mans[p]["products"] for p in positions):
            result["instant"] = finalize_instant(flame, F, fdir, positions)
        else:
            warn(flame, "instantaneous fronts missing for some positions")
    if result["cbar"]:
        x = write_flame_excel(flame, fdir, positions, mans, result["cbar"], result["instant"])
        log(f"    Excel → {x}")
    build_settings_log(flame, fdir, positions, mans)
    atomic_write_json(fres_path, result)
    atomic_write_json(os.path.join(fdir, "flame_done.json"),
                      {"hash": fhash, "finished": datetime.datetime.now().isoformat(timespec="seconds")})
    log(f"[{_ts()}] ✔ {flame} finalised in {time.time() - t0:.0f} s → {fdir}")
    return result


def write_master(flames, results, status):
    if not DO_CBAR:
        log("DO_CBAR is off → no combined c-line Excel.")
        return
    rows = [results_row(f, results[f]["cbar"]) for f in flames
            if results.get(f) and results[f].get("cbar")]
    st_rows = []
    for f in flames:
        r = results.get(f) or {}
        cbs = []
        if r.get("cbar"):
            cbs = [f"{lv}:{d['status']}" for lv, d in r["cbar"]["levels"].items()
                   if d["status"] != "closed"]
        st_rows.append([f, status.get(f, "?"), pos_list_str(flames[f]["positions"]),
                        str(flames[f]["gaps"] or ""), str(r.get("no_bg") or ""),
                        "; ".join(cbs), sum(1 for ff, _ in WARNINGS if ff == f)])
    inst_rows = []
    for f in flames:
        r = results.get(f) or {}
        vals = [i["A_inst_mm2"] for i in (r.get("instant") or []) if "A_inst_mm2" in i]
        if vals:
            a_real = ([round(r["cbar"]["levels"][str(lv)]["A_real"], 2)
                       for lv in sorted(CBAR_LEVELS)] if r.get("cbar") else [])
            inst_rows.append([f, round(float(np.mean(vals)), 2),
                              round(float(np.std(vals)), 2), len(vals)] + vals
                             + [""] * (N_INSTANT_COMPOSITES - len(vals)) + a_real)
    extra = []
    if inst_rows:
        extra = [("Instantaneous A", ["Condition", "A_inst mean [mm²]", "A_inst std [mm²]", "n"]
                  + [f"A_inst #{i} [mm²]" for i in range(1, N_INSTANT_COMPOSITES + 1)]
                  + [f"A_real (c̄={lv}) [mm²]" for lv in sorted(CBAR_LEVELS)], inst_rows)]
    x = write_table(os.path.join(OUTDIR, "ALL_conditions_results.xlsx"),
                    [("All conditions", results_header(), rows)] + extra + [
                     ("Run status", ["Condition", "status", "positions", "missing positions",
                                     "positions without bg", "c-lines not closed",
                                     "# warnings this run"], st_rows)])
    log(f"Combined Excel → {x}")


# ══════════════════════════════════════════════════════════════════════
#  Main
# ══════════════════════════════════════════════════════════════════════
def main():
    if not (DO_CBAR or DO_INSTANT):
        print("ERROR: DO_CBAR and DO_INSTANT are both False — nothing to do."); sys.exit(1)
    if not os.path.isdir(INPUT_DIR):
        print(f"ERROR: INPUT_DIR does not exist: {INPUT_DIR}")
        if _CONFIG_LOCAL is None:
            print("       No config_local.py found — copy config_local.example.py to "
                  "config_local.py and set your paths there.")
        sys.exit(1)
    os.makedirs(OUTDIR, exist_ok=True)
    ld = os.path.join(OUTDIR, "logs"); os.makedirs(ld, exist_ok=True)
    stamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    log_path = os.path.join(ld, f"run_{stamp}.log")
    sys.stdout = _Tee(log_path)

    log(f"run_wabel_batch.py {SCRIPT_VERSION} — started {stamp}")
    log(f"config     = {_CONFIG_LOCAL or 'defaults in run_wabel_batch.py (no config_local.py)'}")
    log(f"INPUT_DIR  = {INPUT_DIR}\nOUTDIR     = {OUTDIR}\nlog file   = {log_path}")
    log(f"DO_CBAR={DO_CBAR}  DO_INSTANT={DO_INSTANT}  NUM_WORKERS={NUM_WORKERS}  "
        f"CHUNK_FRAMES={CHUNK_FRAMES}  N_INSTANT={N_INSTANT}")
    cd = os.path.join(OUTDIR, "code_used"); os.makedirs(cd, exist_ok=True)
    shutil.copy2(os.path.abspath(__file__), os.path.join(cd, os.path.basename(__file__)))
    sr = os.path.join(os.path.dirname(os.path.abspath(__file__)), "spe_reader.py")
    if os.path.isfile(sr):
        shutil.copy2(sr, os.path.join(cd, "spe_reader.py"))
    if _CONFIG_LOCAL and os.path.isfile(_CONFIG_LOCAL):
        shutil.copy2(_CONFIG_LOCAL, os.path.join(cd, "config_local.py"))

    # ── discovery ────────────────────────────────────────────────────
    raw, ignored = discover(INPUT_DIR)
    flames = resolve_flames(raw)
    if not flames:
        log("ERROR: no flames found."); sys.exit(1)
    # which flames are selected for this run
    names = list(flames)
    if EXCLUDE_FLAMES:
        names = [f for f in names if not any(t.lower() in f.lower() for t in EXCLUDE_FLAMES)]
    if ONLY_FLAMES:
        names = [f for f in names if any(t.lower() in f.lower() for t in ONLY_FLAMES)]
    if MAX_FLAMES:
        names = names[:int(MAX_FLAMES)]
    if not names:
        log(f"ERROR: no flame left after ONLY_FLAMES={ONLY_FLAMES} / "
            f"EXCLUDE_FLAMES={EXCLUDE_FLAMES} / MAX_FLAMES={MAX_FLAMES}."); sys.exit(1)
    limited = len(names) < len(flames)

    log(f"\n{'=' * 72}\nDISCOVERY: {len(flames)} flames found, "
        f"{len(names)} selected for this run")
    for p in ignored:
        warn("discovery", f"file name not understood (no 'pos<N>') → ignored: {p}")
    total_bytes = 0
    for f, F in flames.items():
        sel = f in names
        if sel:
            for e in F["positions"].values():
                total_bytes += sum(os.path.getsize(p) for p in e["data"])
                if e["bg"]: total_bytes += os.path.getsize(e["bg"])
        log(f"  {'•' if sel else '-'} {f}   D={F['D']} mm   {pos_list_str(F['positions'])}   "
            f"[folder: {', '.join(F['folders'])}]" + ("" if sel else "   (not in this run)"))
        if sel:                      # warnings only for flames that are processed
            for m in F["disc_warnings"]:
                warn(f, m)
    flames = OrderedDict((f, flames[f]) for f in names)
    if limited:
        log(f"\n  >>> LIMITED RUN (ONLY_FLAMES={ONLY_FLAMES}, EXCLUDE_FLAMES={EXCLUDE_FLAMES}, "
            f"MAX_FLAMES={MAX_FLAMES}): processing {len(flames)} flame(s): {', '.join(flames)}")
    log(f"  Total .spe data (selected flames): {total_bytes / 1e9:.1f} GB")

    # ── resume scan ──────────────────────────────────────────────────
    log(f"\n{'=' * 72}\nRESUME SCAN")
    jobs = []; pending = {}; skipped_pos = {}
    for f, F in flames.items():
        fdir = os.path.join(OUTDIR, safe_name(f))
        pending[f] = set(); skipped_pos[f] = set()
        parts = []
        for pos, ent in F["positions"].items():
            if ent["bg"] is None and MISSING_BG_ACTION == "skip":
                skipped_pos[f].add(pos); parts.append(f"pos{pos}:SKIP(no bg)"); continue
            needs, keep, notes = plan_position(fdir, pos, ent)
            for n in notes:
                warn(f, f"pos{pos}: {n}")
            if needs:
                pending[f].add(pos)
                jobs.append({"flame": f, "pos": pos, "spe_paths": ent["data"],
                             "bg_path": ent["bg"], "pos_dir": os.path.join(fdir, f"pos{pos}"),
                             "needs": needs, "keep_products": keep,
                             "sources": source_sig(ent)})
                parts.append(f"pos{pos}:{'+'.join(needs)}")
            else:
                parts.append(f"pos{pos}:done")
        log(f"  {f}: " + "  ".join(parts))
    n_jobs = len(jobs)
    log(f"\n  → {n_jobs} positions to process, "
        f"{sum(len(F['positions']) for F in flames.values()) - n_jobs} already complete")
    if DRY_RUN:
        log("\nDRY_RUN = True → stopping here."); return

    # ── run ──────────────────────────────────────────────────────────
    n_flames = len(flames)
    results, status = {}, {}
    failed_pos = {f: set() for f in flames}
    state = {"flames_done": 0}
    t_start = time.time()

    def try_finalize(f):
        if failed_pos[f] or skipped_pos[f]:
            bad = sorted(failed_pos[f] | skipped_pos[f])
            status[f] = f"NOT finalised (pos {bad} failed/skipped)"
            warn(f, f"flame NOT finalised: positions {bad} failed or were skipped")
        else:
            try:
                results[f] = finalize_flame(f, flames[f]); status[f] = "ok"
            except Exception as e:
                status[f] = f"finalisation failed: {e}"
                warn(f, f"FINALISATION FAILED: {type(e).__name__}: {e}")
                log(traceback.format_exc())
        state["flames_done"] += 1
        log(f"[{_ts()}] ■ flames {state['flames_done']}/{n_flames} complete")

    log(f"\n{'=' * 72}\nPROCESSING  ({n_jobs} positions, {min(NUM_WORKERS, max(n_jobs, 1))} workers)")
    pool = None
    try:
        if n_jobs:
            if NUM_WORKERS > 1 and n_jobs > 1:
                from multiprocessing import Pool
                pool = Pool(min(NUM_WORKERS, n_jobs), initializer=_init_worker)
                it = pool.imap_unordered(process_position, jobs, chunksize=1)
            else:
                _init_worker()
                it = (process_position(j) for j in jobs)
        else:
            it = iter(())
        # flames with nothing to process → finalise while workers run
        for f in flames:
            if not pending[f]:
                try_finalize(f)
        done_jobs = 0
        for r in it:
            done_jobs += 1
            f, p = r["flame"], r["pos"]
            for w in r["warnings"]:
                warn(f, f"pos{p}: {w}")
            pending[f].discard(p)
            ntot = len(flames[f]["positions"])
            ndone = ntot - len(pending[f]) - len(skipped_pos[f])
            el = time.time() - t_start
            eta = el / done_jobs * (n_jobs - done_jobs)
            if r["ok"]:
                mark = f"✓ {f}  pos{p} [{'+'.join(r['needs'])}] {r['n_frames']} fr, {r['seconds']:.0f} s"
            else:
                failed_pos[f].add(p)
                mark = f"✗ {f}  pos{p} FAILED after {r['seconds']:.0f} s"
                warn(f, f"pos{p}: PROCESSING FAILED: {r['error']}")
            log(f"[{_ts()}] {mark}\n"
                f"           flame {ndone}/{ntot} positions | all positions "
                f"{done_jobs}/{n_jobs} | flames {state['flames_done']}/{n_flames} | "
                f"elapsed {fmt_dur(el)} | ETA {fmt_dur(eta)}")
            if not pending[f]:
                try_finalize(f)
    except KeyboardInterrupt:
        log("\n!!! Interrupted — completed positions are saved; rerun to resume.")
        if pool: pool.terminate()
        raise
    finally:
        if pool:
            pool.close(); pool.join()

    write_master(flames, results, status)

    log(f"\n{'=' * 72}\nSUMMARY  (total {fmt_dur(time.time() - t_start)})")
    for f in flames:
        log(f"  {f}: {status.get(f, '?')}")
    if WARNINGS:
        log(f"\n{len(WARNINGS)} warnings (also in {log_path}):")
        for f, m in WARNINGS:
            log(f"  [{f}] {m.splitlines()[0]}")
    log("\nDone.")


if __name__ == "__main__":
    main()
